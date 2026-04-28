"""Document text extraction. Supports .docx, .doc (via pre-extracted .txt sidecar), .pdf, .txt."""
from pathlib import Path
from docx import Document as DocxDocument
from pypdf import PdfReader


def extract_text(path: Path) -> str:
    """Best-effort plain text extraction. Falls back to a sidecar .txt if .doc isn't readable."""
    suffix = path.suffix.lower()

    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")

    if suffix == ".docx":
        try:
            d = DocxDocument(str(path))
            parts: list[str] = []
            for p in d.paragraphs:
                parts.append(p.text)
            # Tables — flatten to tab-separated rows
            for t in d.tables:
                for row in t.rows:
                    parts.append("\t".join(cell.text for cell in row.cells))
            return "\n".join(parts)
        except Exception as e:  # noqa: BLE001
            return _fallback_sidecar(path) or f"[docx extract error: {e}]"

    if suffix == ".doc":
        return _fallback_sidecar(path) or f"[no extractor for legacy .doc: {path.name}]"

    if suffix == ".pdf":
        try:
            r = PdfReader(str(path))
            parts = [page.extract_text() or "" for page in r.pages]
            return "\n".join(parts)
        except Exception as e:  # noqa: BLE001
            return f"[pdf extract error: {e}]"

    return path.read_text(encoding="utf-8", errors="replace")


def _fallback_sidecar(path: Path) -> str | None:
    """If a sibling `<file>.<ext>.txt` exists (textutil pre-extracted), use it."""
    sidecar = path.with_suffix(path.suffix + ".txt")
    if sidecar.exists():
        return sidecar.read_text(encoding="utf-8", errors="replace")
    return None
