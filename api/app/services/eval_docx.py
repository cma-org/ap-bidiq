"""Generate the Evaluation Statement as a DOCX file matching AP's human format."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor
from sqlalchemy.orm import Session

from app.models import Bid, EvalStatement, Tender


def _shade_cell(cell, hex_color: str) -> None:
    """Apply background fill to a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def render_eval_statement_docx(db: Session, bid_id: int) -> bytes:
    es = (
        db.query(EvalStatement)
        .filter(EvalStatement.bid_id == bid_id)
        .order_by(EvalStatement.generated_at.desc())
        .first()
    )
    if not es:
        raise ValueError(f"No evaluation statement for bid {bid_id}")

    bid = db.get(Bid, bid_id)
    tender = db.get(Tender, bid.tender_id)
    payload = es.json_payload or {}

    doc = Document()

    # ---- Header / metadata
    section = doc.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("EVALUATION STATEMENT")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x0C, 0x4A, 0x6E)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("Government of Andhra Pradesh — Infrastructure & Investment Department")
    sr.italic = True
    sr.font.size = Pt(10)
    sr.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    doc.add_paragraph()

    # Metadata block as a 2-col table
    meta = doc.add_table(rows=5, cols=2)
    meta.autofit = True
    meta_data = [
        ("Tender:", tender.title),
        ("Tender code:", tender.code or "—"),
        ("Vendor:", bid.vendor_name),
        ("Bid value:", f"₹{bid.bid_value_inr:,.2f} Cr" if bid.bid_value_inr else "—"),
        ("Active rules version:", f"v{payload.get('active_rules_version', '?')}"),
    ]
    for i, (label, value) in enumerate(meta_data):
        c0 = meta.cell(i, 0)
        c1 = meta.cell(i, 1)
        c0.width = Cm(5)
        run = c0.paragraphs[0].add_run(label)
        run.bold = True
        run.font.size = Pt(10)
        c1.paragraphs[0].add_run(str(value)).font.size = Pt(10)

    doc.add_paragraph()

    # ---- Verdict banner
    verdict = es.verdict
    verdict_label = {"qualified": "QUALIFIED", "not_qualified": "NOT QUALIFIED", "conditional": "CONDITIONAL"}[verdict]
    verdict_color = {"qualified": "047857", "not_qualified": "B91C1C", "conditional": "B45309"}[verdict]

    vbar = doc.add_table(rows=1, cols=1)
    vcell = vbar.cell(0, 0)
    _shade_cell(vcell, verdict_color)
    vp = vcell.paragraphs[0]
    vp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    vrun = vp.add_run(f"VERDICT: {verdict_label}")
    vrun.bold = True
    vrun.font.size = Pt(16)
    vrun.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    if es.summary:
        sp = doc.add_paragraph()
        sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sr = sp.add_run(es.summary)
        sr.italic = True
        sr.font.size = Pt(10)
        sr.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    doc.add_paragraph()

    # ---- Evaluation matrix table
    rows = payload.get("rows", [])

    matrix = doc.add_table(rows=1 + len(rows), cols=6)
    matrix.style = "Light Grid Accent 1"

    headers = ["Check ID", "Criterion", "Required", "Submitted", "Meets?", "Remarks"]
    header_row = matrix.rows[0]
    for i, h in enumerate(headers):
        cell = header_row.cells[i]
        _shade_cell(cell, "0C4A6E")
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(0xF8, 0xFA, 0xFC)
        run.font.size = Pt(9)

    for r_idx, r in enumerate(rows, start=1):
        row = matrix.rows[r_idx]
        if r["verdict"] == "fail":
            for c in row.cells:
                _shade_cell(c, "FEE2E2")
        cells_data = [
            r.get("check_id", ""),
            r.get("criterion", ""),
            r.get("required", ""),
            r.get("submitted", ""),
            "✓" if r.get("meets") else "✗",
            r.get("remarks", ""),
        ]
        for i, val in enumerate(cells_data):
            cell = row.cells[i]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            run = cell.paragraphs[0].add_run(str(val))
            run.font.size = Pt(8)
            if i == 4:  # the meets column
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                run.bold = True
                if r.get("meets"):
                    run.font.color.rgb = RGBColor(0x04, 0x78, 0x57)
                else:
                    run.font.color.rgb = RGBColor(0xB9, 0x1C, 0x1C)

    doc.add_paragraph()

    # ---- Reasons (only for non-qualified)
    if es.reasons:
        h = doc.add_paragraph()
        hr = h.add_run("Disqualification reasons:")
        hr.bold = True
        hr.font.size = Pt(11)
        hr.font.color.rgb = RGBColor(0xB9, 0x1C, 0x1C)

        for r in es.reasons:
            p = doc.add_paragraph(style="List Number")
            t_run = p.add_run(f"{r['title']} — ")
            t_run.bold = True
            t_run.font.size = Pt(10)
            p.add_run(r["message"]).font.size = Pt(10)

        doc.add_paragraph()

    # ---- Footer / signoff
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run(
        "Generated by AP-BidIQ · Cited from official tender + Corrigendum-1 · "
        "AI verdict for officer review; final decision rests with the Procurement Officer."
    )
    fr.italic = True
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    sign = doc.add_paragraph()
    sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    sign.add_run("\n\n").font.size = Pt(10)
    sign_run = sign.add_run("Procurement Officer:  _____________________________   Date: _____________")
    sign_run.font.size = Pt(10)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
